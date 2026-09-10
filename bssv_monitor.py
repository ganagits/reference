#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
BSSV Monitor
============

Posts SOAP payloads to one or more Oracle JD Edwards BSSV endpoints, decides
success/failure, logs every check to a rolling CSV, republishes a 72-hour JSON
feed for the dashboard, and raises e-mail alerts (via PowerShell
Send-MailMessage) only when a check fails.

Usage
-----
    python bssv_monitor.py                 # use run_mode from config.ini
    python bssv_monitor.py --once          # single pass (Task Scheduler)
    python bssv_monitor.py --loop          # stay resident
    python bssv_monitor.py --config D:\\cfg\\config.ini
    python bssv_monitor.py --once --dry-run    # probe, but never send mail
    python bssv_monitor.py --test-email        # send a test alert and exit

Author: generated for the BSSV monitoring project.
"""

from __future__ import annotations

import argparse
import configparser
import csv
import html
import json
import logging
import logging.handlers
import os
import platform
import re
import shutil
import subprocess
import sys
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone

try:
    import requests
    from requests.adapters import HTTPAdapter
except ImportError:  # pragma: no cover
    sys.exit("Missing dependency 'requests'.  Run:  pip install -r requirements.txt")

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    sys.exit("Missing dependency 'openpyxl'.  Run:  pip install -r requirements.txt")

import urllib3

APP_NAME = "BSSV Monitor"
APP_VERSION = "1.0.0"
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

CSV_COLUMNS = [
    "run_id",
    "timestamp_local",
    "timestamp_utc",
    "server_name",
    "environment",
    "service_name",
    "endpoint_url",
    "payload_file",
    "status",
    "http_status",
    "response_ms",
    "failure_reason",
    "detail",
    "attempts",
    "email_sent",
]

log = logging.getLogger("bssv")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def resolve(path: str, base: str = BASE_DIR) -> str:
    """Absolute path; relative values are resolved against the project folder."""
    if not path:
        return ""
    path = os.path.expandvars(os.path.expanduser(path.strip().strip('"')))
    if os.sep != "\\":  # allow the Windows-style config to be tested on POSIX
        path = path.replace("\\", "/")
    return path if os.path.isabs(path) else os.path.normpath(os.path.join(base, path))


def utcnow() -> datetime:
    return datetime.now(timezone.utc)


def iso_utc(dt: datetime) -> str:
    return dt.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def iso_local(dt: datetime) -> str:
    return dt.astimezone().strftime("%Y-%m-%d %H:%M:%S")


def parse_utc(text: str):
    """Parse a timestamp_utc value back into a datetime (None if unparseable)."""
    if not text:
        return None
    text = text.strip()
    for fmt in ("%Y-%m-%dT%H:%M:%SZ", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    try:
        dt = datetime.fromisoformat(text.replace("Z", "+00:00"))
        return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)
    except ValueError:
        return None


def as_bool(value, default=False) -> bool:
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in ("1", "true", "yes", "y", "on", "enabled"):
        return True
    if text in ("0", "false", "no", "n", "off", "disabled"):
        return False
    return default


def as_int(value, default):
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return default


def slugify(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", str(text or "unknown")).strip("_") or "unknown"


# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
class Config:
    def __init__(self, path: str):
        self.path = path
        parser = configparser.ConfigParser(inline_comment_prefixes=(";",))
        parser.optionxform = str.lower
        if not os.path.isfile(path):
            sys.exit(f"Config file not found: {path}")
        parser.read(path, encoding="utf-8-sig")
        self.cp = parser
        base = os.path.dirname(path)

        g = self._sec("general")
        self.run_interval_seconds = as_int(g.get("run_interval_seconds"), 300)
        self.run_mode = (g.get("run_mode") or "loop").strip().lower()
        self.skip_missed_ticks = as_bool(g.get("skip_missed_ticks"), True)
        self.excel_config = resolve(g.get("excel_config") or "servers.xlsx", base)
        self.excel_sheet = (g.get("excel_sheet") or "Servers").strip()
        self.payload_dir = resolve(g.get("payload_dir") or "payloads", base)
        self.max_workers = max(1, as_int(g.get("max_workers"), 4))

        r = self._sec("request")
        self.timeout_seconds = as_int(r.get("timeout_seconds"), 30)
        self.retries = max(0, as_int(r.get("retries"), 1))
        self.retry_delay_seconds = as_int(r.get("retry_delay_seconds"), 5)
        self.verify_ssl = as_bool(r.get("verify_ssl"), False)
        self.ca_bundle = resolve(r.get("ca_bundle") or "", base)
        self.default_soap_action = (r.get("default_soap_action") or '""').strip()
        self.content_type = (r.get("content_type") or "text/xml; charset=utf-8").strip()
        self.user_agent = (r.get("user_agent") or f"{APP_NAME}/{APP_VERSION}").strip()
        self.http_proxy = (r.get("http_proxy") or "").strip()
        self.https_proxy = (r.get("https_proxy") or "").strip()

        lg = self._sec("logging")
        self.result_log = resolve(lg.get("result_log") or "logs/bssv_monitor_log.csv", base)
        self.retention_hours = as_int(lg.get("retention_hours"), 72)
        self.archive_purged = as_bool(lg.get("archive_purged"), True)
        self.archive_dir = resolve(lg.get("archive_dir") or "logs/archive", base)
        self.app_log = resolve(lg.get("app_log") or "logs/bssv_monitor.log", base)
        self.app_log_level = (lg.get("app_log_level") or "INFO").strip().upper()
        self.app_log_max_bytes = as_int(lg.get("app_log_max_bytes"), 5 * 1024 * 1024)
        self.app_log_backup_count = as_int(lg.get("app_log_backup_count"), 5)
        self.save_failed_responses = as_bool(lg.get("save_failed_responses"), True)
        self.failed_response_dir = resolve(lg.get("failed_response_dir") or "logs/responses", base)
        self.failed_response_retention_hours = as_int(
            lg.get("failed_response_retention_hours"), 72
        )

        d = self._sec("dashboard")
        self.dashboard_enabled = as_bool(d.get("enabled"), True)
        self.web_root = resolve(d.get("web_root") or "dashboard", base)
        self.json_file_name = (d.get("json_file_name") or "status.json").strip()
        self.window_hours = as_int(d.get("window_hours"), 72)
        self.copy_dashboard_files = as_bool(d.get("copy_dashboard_files"), True)
        self.max_detail_rows = as_int(d.get("max_detail_rows"), 3000)

        e = self._sec("email")
        self.email_enabled = as_bool(e.get("enabled"), True)
        self.powershell_script = resolve(e.get("powershell_script") or "send_alert.ps1", base)
        self.powershell_exe = (e.get("powershell_exe") or "powershell.exe").strip()
        self.smtp_server = (e.get("smtp_server") or "").strip()
        self.smtp_port = as_int(e.get("smtp_port"), 25)
        self.use_ssl = as_bool(e.get("use_ssl"), False)
        self.from_address = (e.get("from_address") or "").strip()
        self.to_addresses = (e.get("to_addresses") or "").strip()
        self.cc_addresses = (e.get("cc_addresses") or "").strip()
        self.subject_prefix = (e.get("subject_prefix") or "[BSSV ALERT]").strip()
        self.throttle_minutes = as_int(e.get("throttle_minutes"), 30)
        self.consolidate = as_bool(e.get("consolidate"), True)
        self.group_by_environment = as_bool(e.get("group_by_environment"), True)
        # [email.recipients] - per-environment routing, e.g.  PROD = ops@x.com
        self.env_recipients = {
            str(k).strip().upper(): str(v).strip()
            for k, v in self._sec("email.recipients").items()
            if str(v).strip()
        }
        self.attach_response = as_bool(e.get("attach_response"), True)
        self.send_recovery = as_bool(e.get("send_recovery"), False)

        self.alert_state_file = os.path.join(
            os.path.dirname(self.result_log) or base, "alert_state.json"
        )

    def _sec(self, name: str) -> dict:
        return dict(self.cp[name]) if self.cp.has_section(name) else {}


def setup_logging(cfg: Config, verbose: bool = False) -> None:
    os.makedirs(os.path.dirname(cfg.app_log) or ".", exist_ok=True)
    level = getattr(logging, cfg.app_log_level, logging.INFO)
    log.setLevel(logging.DEBUG if verbose else level)
    log.handlers.clear()
    fmt = logging.Formatter("%(asctime)s  %(levelname)-7s  %(message)s", "%Y-%m-%d %H:%M:%S")

    fh = logging.handlers.RotatingFileHandler(
        cfg.app_log,
        maxBytes=cfg.app_log_max_bytes,
        backupCount=cfg.app_log_backup_count,
        encoding="utf-8",
    )
    fh.setFormatter(fmt)
    log.addHandler(fh)

    sh = logging.StreamHandler(sys.stdout)
    sh.setFormatter(fmt)
    log.addHandler(sh)


# ---------------------------------------------------------------------------
# Excel server list
# ---------------------------------------------------------------------------
HEADER_ALIASES = {
    "enabled": "enabled",
    "servername": "server_name",
    "server": "server_name",
    "environment": "environment",
    "env": "environment",
    "servicename": "service_name",
    "service": "service_name",
    "endpointurl": "endpoint_url",
    "endpoint": "endpoint_url",
    "url": "endpoint_url",
    "payloadpath": "payload_path",
    "payload": "payload_path",
    "payloadfile": "payload_path",
    "soapaction": "soap_action",
    "expectedtext": "expected_text",
    "expected": "expected_text",
    "timeoutseconds": "timeout_seconds",
    "timeout": "timeout_seconds",
    "verifyssl": "verify_ssl",
    "alertemailto": "alert_email_to",
    "emailto": "alert_email_to",
    "notes": "notes",
}


def read_servers(cfg: Config) -> list:
    """Read the Excel workbook into a list of server definition dicts."""
    if not os.path.isfile(cfg.excel_config):
        raise FileNotFoundError(f"Excel config not found: {cfg.excel_config}")

    wb = load_workbook(cfg.excel_config, data_only=True, read_only=True)
    sheet = wb[cfg.excel_sheet] if cfg.excel_sheet in wb.sheetnames else wb[wb.sheetnames[0]]

    rows = list(sheet.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []

    header_map = {}
    for idx, cell in enumerate(rows[0]):
        key = re.sub(r"[^a-z0-9]", "", str(cell or "").lower())
        if key in HEADER_ALIASES:
            header_map[HEADER_ALIASES[key]] = idx

    required = {"server_name", "environment", "endpoint_url", "payload_path"}
    missing = required - set(header_map)
    if missing:
        raise ValueError(
            f"Excel sheet '{sheet.title}' is missing required column(s): "
            + ", ".join(sorted(missing))
        )

    def cell(row, name, default=""):
        idx = header_map.get(name)
        if idx is None or idx >= len(row):
            return default
        value = row[idx]
        return default if value is None else str(value).strip()

    servers = []
    for line_no, row in enumerate(rows[1:], start=2):
        if row is None or all(v in (None, "") for v in row):
            continue
        if not as_bool(cell(row, "enabled", "Y"), True):
            continue

        name = cell(row, "server_name")
        env = cell(row, "environment").upper()
        url = cell(row, "endpoint_url")
        payload = cell(row, "payload_path")
        if not (name and env and url and payload):
            log.warning(
                "Excel row %s skipped - ServerName/Environment/EndpointURL/PayloadPath incomplete",
                line_no,
            )
            continue

        servers.append(
            {
                "row": line_no,
                "server_name": name,
                "environment": env,
                "service_name": cell(row, "service_name", "-") or "-",
                "endpoint_url": url,
                "payload_path": resolve(payload, cfg.payload_dir),
                "soap_action": cell(row, "soap_action", cfg.default_soap_action),
                "expected_text": cell(row, "expected_text"),
                "timeout_seconds": as_int(cell(row, "timeout_seconds"), cfg.timeout_seconds),
                "verify_ssl": as_bool(cell(row, "verify_ssl"), cfg.verify_ssl),
                "alert_email_to": (
                    cell(row, "alert_email_to")            # 1. the Excel cell
                    or cfg.env_recipients.get(env, "")     # 2. [email.recipients] for this env
                    or cfg.to_addresses                    # 3. the global default
                ),
                "notes": cell(row, "notes"),
            }
        )
    return servers


# ---------------------------------------------------------------------------
# Probing
# ---------------------------------------------------------------------------
FAULT_RE = re.compile(r"<\s*(?:[a-zA-Z0-9_.-]+:)?Fault[\s>]", re.IGNORECASE)
FAULTSTRING_RE = re.compile(
    r"<\s*(?:[a-zA-Z0-9_.-]+:)?(?:faultstring|Text|message)\s*[^>]*>(.*?)<\s*/",
    re.IGNORECASE | re.DOTALL,
)


def build_session(cfg: Config) -> requests.Session:
    session = requests.Session()
    adapter = HTTPAdapter(pool_connections=cfg.max_workers * 2, pool_maxsize=cfg.max_workers * 2)
    session.mount("https://", adapter)
    session.mount("http://", adapter)
    proxies = {}
    if cfg.http_proxy:
        proxies["http"] = cfg.http_proxy
    if cfg.https_proxy:
        proxies["https"] = cfg.https_proxy
    if proxies:
        session.proxies.update(proxies)
    elif cfg.http_proxy == "" and cfg.https_proxy == "":
        session.trust_env = True
    return session


def evaluate(server: dict, response, body: str) -> tuple:
    """Return (ok, failure_reason, detail).

    Success rule: HTTP 200  AND  no SOAP Fault  AND  ExpectedText present
    (the ExpectedText check is skipped when the Excel cell is blank).
    """
    if response.status_code != 200:
        return False, "HTTP_STATUS", f"HTTP {response.status_code} {response.reason}"

    if FAULT_RE.search(body):
        match = FAULTSTRING_RE.search(body)
        fault_text = " ".join(match.group(1).split())[:400] if match else "SOAP Fault returned"
        return False, "SOAP_FAULT", fault_text

    expected = server.get("expected_text") or ""
    if expected and expected not in body:
        return False, "EXPECTED_TEXT_MISSING", f"Response did not contain: {expected[:200]}"

    return True, "", ""


def probe(cfg: Config, session: requests.Session, server: dict, run_id: str) -> dict:
    started = utcnow()
    result = {
        "run_id": run_id,
        "timestamp_local": iso_local(started),
        "timestamp_utc": iso_utc(started),
        "server_name": server["server_name"],
        "environment": server["environment"],
        "service_name": server["service_name"],
        "endpoint_url": server["endpoint_url"],
        "payload_file": os.path.basename(server["payload_path"]),
        "status": "FAILED",
        "http_status": "",
        "response_ms": "",
        "failure_reason": "",
        "detail": "",
        "attempts": 0,
        "email_sent": "NO",
        "_body": "",
    }

    try:
        with open(server["payload_path"], "rb") as handle:
            payload = handle.read()
    except OSError as exc:
        result["failure_reason"] = "PAYLOAD_UNREADABLE"
        result["detail"] = str(exc)[:400]
        result["attempts"] = 0
        return result

    verify = server["verify_ssl"]
    if verify and cfg.ca_bundle and os.path.isfile(cfg.ca_bundle):
        verify = cfg.ca_bundle

    soap_action = server.get("soap_action") or '""'
    headers = {
        "Content-Type": cfg.content_type,
        "SOAPAction": soap_action,
        "User-Agent": cfg.user_agent,
        "Accept": "text/xml, application/soap+xml, */*",
        "Connection": "close",
    }

    total_attempts = cfg.retries + 1
    for attempt in range(1, total_attempts + 1):
        result["attempts"] = attempt
        clock = time.perf_counter()
        try:
            response = session.post(
                server["endpoint_url"],
                data=payload,
                headers=headers,
                timeout=server["timeout_seconds"],
                verify=verify,
            )
            elapsed = int((time.perf_counter() - clock) * 1000)
            body = response.text or ""
            result["http_status"] = str(response.status_code)
            result["response_ms"] = str(elapsed)
            result["_body"] = body

            ok, reason, detail = evaluate(server, response, body)
            if ok:
                result["status"] = "SUCCESS"
                result["failure_reason"] = ""
                result["detail"] = ""
                result["_body"] = ""
                return result

            result["failure_reason"] = reason
            result["detail"] = detail

        except requests.exceptions.SSLError as exc:
            result["response_ms"] = str(int((time.perf_counter() - clock) * 1000))
            result["failure_reason"] = "SSL_ERROR"
            result["detail"] = str(exc)[:400]
        except requests.exceptions.ConnectTimeout as exc:
            result["response_ms"] = str(int((time.perf_counter() - clock) * 1000))
            result["failure_reason"] = "CONNECT_TIMEOUT"
            result["detail"] = str(exc)[:400]
        except requests.exceptions.ReadTimeout as exc:
            result["response_ms"] = str(int((time.perf_counter() - clock) * 1000))
            result["failure_reason"] = "READ_TIMEOUT"
            result["detail"] = str(exc)[:400]
        except requests.exceptions.ConnectionError as exc:
            result["response_ms"] = str(int((time.perf_counter() - clock) * 1000))
            result["failure_reason"] = "CONNECTION_ERROR"
            result["detail"] = str(exc)[:400]
        except Exception as exc:  # noqa: BLE001 - never let one server kill the pass
            result["response_ms"] = str(int((time.perf_counter() - clock) * 1000))
            result["failure_reason"] = "UNEXPECTED_ERROR"
            result["detail"] = f"{type(exc).__name__}: {exc}"[:400]

        if attempt < total_attempts:
            log.debug(
                "%s attempt %s/%s failed (%s) - retrying in %ss",
                server["server_name"], attempt, total_attempts,
                result["failure_reason"], cfg.retry_delay_seconds,
            )
            time.sleep(max(0, cfg.retry_delay_seconds))

    return result


def run_pass(cfg: Config, servers: list, run_id: str) -> list:
    session = build_session(cfg)
    try:
        if cfg.max_workers == 1:
            return [probe(cfg, session, s, run_id) for s in servers]
        with ThreadPoolExecutor(max_workers=cfg.max_workers) as pool:
            return list(pool.map(lambda s: probe(cfg, session, s, run_id), servers))
    finally:
        session.close()


# ---------------------------------------------------------------------------
# CSV log + autopurge
# ---------------------------------------------------------------------------
def append_results(cfg: Config, results: list) -> None:
    os.makedirs(os.path.dirname(cfg.result_log) or ".", exist_ok=True)
    new_file = not os.path.isfile(cfg.result_log) or os.path.getsize(cfg.result_log) == 0
    with open(cfg.result_log, "a", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=CSV_COLUMNS, extrasaction="ignore")
        if new_file:
            writer.writeheader()
        for row in results:
            writer.writerow({k: row.get(k, "") for k in CSV_COLUMNS})


def read_log(cfg: Config) -> list:
    if not os.path.isfile(cfg.result_log):
        return []
    with open(cfg.result_log, "r", newline="", encoding="utf-8-sig") as handle:
        return [dict(r) for r in csv.DictReader(handle)]


def purge_log(cfg: Config) -> int:
    """Drop rows older than retention_hours; optionally archive them first."""
    rows = read_log(cfg)
    if not rows:
        return 0

    cutoff = utcnow() - timedelta(hours=cfg.retention_hours)
    keep, drop = [], []
    for row in rows:
        stamp = parse_utc(row.get("timestamp_utc", ""))
        (keep if (stamp is None or stamp >= cutoff) else drop).append(row)

    if not drop:
        return 0

    if cfg.archive_purged:
        os.makedirs(cfg.archive_dir, exist_ok=True)
        buckets = {}
        for row in drop:
            stamp = parse_utc(row.get("timestamp_utc", "")) or utcnow()
            buckets.setdefault(stamp.strftime("%Y%m"), []).append(row)
        for bucket, bucket_rows in buckets.items():
            target = os.path.join(cfg.archive_dir, f"bssv_monitor_log_{bucket}.csv")
            new_file = not os.path.isfile(target) or os.path.getsize(target) == 0
            with open(target, "a", newline="", encoding="utf-8-sig") as handle:
                writer = csv.DictWriter(handle, fieldnames=CSV_COLUMNS, extrasaction="ignore")
                if new_file:
                    writer.writeheader()
                for row in bucket_rows:
                    writer.writerow({k: row.get(k, "") for k in CSV_COLUMNS})

    tmp = cfg.result_log + ".tmp"
    with open(tmp, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=CSV_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        for row in keep:
            writer.writerow({k: row.get(k, "") for k in CSV_COLUMNS})
    os.replace(tmp, cfg.result_log)

    log.info("Autopurge: removed %s row(s) older than %sh", len(drop), cfg.retention_hours)
    return len(drop)


def purge_response_files(cfg: Config) -> None:
    if not (cfg.save_failed_responses and os.path.isdir(cfg.failed_response_dir)):
        return
    cutoff = time.time() - cfg.failed_response_retention_hours * 3600
    removed = 0
    for name in os.listdir(cfg.failed_response_dir):
        path = os.path.join(cfg.failed_response_dir, name)
        try:
            if os.path.isfile(path) and os.path.getmtime(path) < cutoff:
                os.remove(path)
                removed += 1
        except OSError:
            pass
    if removed:
        log.info("Autopurge: removed %s stale response file(s)", removed)


def save_failed_response(cfg: Config, result: dict) -> str:
    if not cfg.save_failed_responses or not result.get("_body"):
        return ""
    os.makedirs(cfg.failed_response_dir, exist_ok=True)
    name = "{}_{}_{}.xml".format(
        utcnow().strftime("%Y%m%d_%H%M%S"),
        slugify(result["server_name"]),
        slugify(result["service_name"]),
    )
    path = os.path.join(cfg.failed_response_dir, name)
    try:
        with open(path, "w", encoding="utf-8", errors="replace") as handle:
            handle.write(result["_body"])
        return path
    except OSError as exc:
        log.warning("Could not save failed response: %s", exc)
        return ""


# ---------------------------------------------------------------------------
# status.json for the dashboard
# ---------------------------------------------------------------------------
def build_status_json(cfg: Config, rows: list) -> dict:
    now = utcnow()
    cutoff = now - timedelta(hours=cfg.window_hours)

    window = []
    for row in rows:
        stamp = parse_utc(row.get("timestamp_utc", ""))
        if stamp and stamp >= cutoff:
            row["_ts"] = stamp
            window.append(row)
    window.sort(key=lambda r: r["_ts"], reverse=True)  # newest first

    success = sum(1 for r in window if r.get("status") == "SUCCESS")
    failed = len(window) - success

    # ---- per-server rollup -------------------------------------------------
    servers = {}
    for row in window:  # newest first, so the first hit is the latest check
        key = (row.get("server_name", ""), row.get("environment", ""), row.get("service_name", ""))
        entry = servers.get(key)
        if entry is None:
            entry = servers[key] = {
                "server_name": key[0],
                "environment": key[1],
                "service_name": key[2],
                "endpoint_url": row.get("endpoint_url", ""),
                "last_status": row.get("status", ""),
                "last_checked_utc": row.get("timestamp_utc", ""),
                "last_checked_local": row.get("timestamp_local", ""),
                "last_failure_reason": "",
                "last_detail": "",
                "total": 0,
                "success": 0,
                "failed": 0,
                "avg_response_ms": 0,
                "uptime_pct": 0.0,
                "_ms": [],
            }
            if row.get("status") != "SUCCESS":
                entry["last_failure_reason"] = row.get("failure_reason", "")
                entry["last_detail"] = row.get("detail", "")
        entry["total"] += 1
        if row.get("status") == "SUCCESS":
            entry["success"] += 1
        else:
            entry["failed"] += 1
        ms = as_int(row.get("response_ms"), 0)
        if ms:
            entry["_ms"].append(ms)

    server_list = []
    for entry in servers.values():
        samples = entry.pop("_ms")
        entry["avg_response_ms"] = int(sum(samples) / len(samples)) if samples else 0
        entry["uptime_pct"] = round(entry["success"] * 100.0 / entry["total"], 2) if entry["total"] else 0.0
        server_list.append(entry)
    server_list.sort(key=lambda e: (e["last_status"] == "SUCCESS", e["server_name"].lower()))

    # ---- per-environment rollup --------------------------------------------
    env_stats = {}
    for entry in server_list:
        env = entry["environment"] or "-"
        stat = env_stats.setdefault(
            env,
            {"name": env, "all": 0, "success": 0, "failed": 0, "success_rate": 0.0,
             "servers_total": 0, "servers_down": 0},
        )
        stat["all"] += entry["total"]
        stat["success"] += entry["success"]
        stat["failed"] += entry["failed"]
        stat["servers_total"] += 1
        if entry["last_status"] != "SUCCESS":
            stat["servers_down"] += 1
    for stat in env_stats.values():
        stat["success_rate"] = round(stat["success"] * 100.0 / stat["all"], 2) if stat["all"] else 0.0

    ranking = ["PROD", "PRODUCTION", "PRD", "LIVE", "PRE-PROD", "PREPROD", "DR",
               "UAT", "TEST", "QA", "DEV", "TRAINING"]

    def env_key(name):
        upper = (name or "").upper()
        return (ranking.index(upper) if upper in ranking else len(ranking), upper)

    environments = sorted(env_stats.values(), key=lambda s: env_key(s["name"]))

    # Group the server cards by environment (PROD first), down servers on top.
    server_list.sort(
        key=lambda e: (env_key(e["environment"]), e["last_status"] == "SUCCESS",
                       e["server_name"].lower())
    )

    # ---- hourly buckets (oldest -> newest, for the chart) -------------------
    # Each bucket carries the overall counts plus a per-environment breakdown,
    # so the dashboard's environment selector can redraw the chart client-side.
    env_names = [s["name"] for s in environments]
    top_of_hour = now.replace(minute=0, second=0, microsecond=0)
    buckets = {}
    for offset in range(cfg.window_hours - 1, -1, -1):
        stamp = top_of_hour - timedelta(hours=offset)
        buckets[stamp.strftime("%Y-%m-%dT%H:00Z")] = {
            "hour_utc": stamp.strftime("%Y-%m-%dT%H:00Z"),
            "hour_local": stamp.astimezone().strftime("%d %b %H:00"),
            "success": 0,
            "failed": 0,
            "by_env": {name: {"success": 0, "failed": 0} for name in env_names},
        }
    for row in window:
        key = row["_ts"].strftime("%Y-%m-%dT%H:00Z")
        bucket = buckets.get(key)
        if not bucket:
            continue
        field = "success" if row.get("status") == "SUCCESS" else "failed"
        bucket[field] += 1
        env_bucket = bucket["by_env"].get(row.get("environment") or "-")
        if env_bucket is not None:
            env_bucket[field] += 1

    # ---- individual checks, newest first -----------------------------------
    checks = []
    for row in window[: cfg.max_detail_rows]:
        checks.append(
            {
                "timestamp_utc": row.get("timestamp_utc", ""),
                "timestamp_local": row.get("timestamp_local", ""),
                "server_name": row.get("server_name", ""),
                "environment": row.get("environment", ""),
                "service_name": row.get("service_name", ""),
                "endpoint_url": row.get("endpoint_url", ""),
                "status": row.get("status", ""),
                "http_status": row.get("http_status", ""),
                "response_ms": as_int(row.get("response_ms"), 0),
                "failure_reason": row.get("failure_reason", ""),
                "detail": row.get("detail", ""),
                "attempts": as_int(row.get("attempts"), 0),
            }
        )

    for row in window:
        row.pop("_ts", None)

    return {
        "generated_at_utc": iso_utc(now),
        "generated_at_local": iso_local(now),
        "window_hours": cfg.window_hours,
        "refresh_seconds": 60,
        "run_interval_seconds": cfg.run_interval_seconds,
        "summary": {
            "all": len(window),
            "success": success,
            "failed": failed,
            "success_rate": round(success * 100.0 / len(window), 2) if window else 0.0,
            "servers_total": len(server_list),
            "servers_down": sum(1 for s in server_list if s["last_status"] != "SUCCESS"),
            "environments_total": len(environments),
            "environments_affected": sum(1 for s in environments if s["servers_down"]),
        },
        "environments": environments,
        "servers": server_list,
        "hourly": list(buckets.values()),
        "checks": checks,
        "truncated": len(window) > cfg.max_detail_rows,
    }


def publish_dashboard(cfg: Config, payload: dict) -> None:
    if not cfg.dashboard_enabled:
        return
    try:
        os.makedirs(cfg.web_root, exist_ok=True)
        target = os.path.join(cfg.web_root, cfg.json_file_name)
        tmp = target + ".tmp"
        with open(tmp, "w", encoding="utf-8") as handle:
            json.dump(payload, handle, indent=1, ensure_ascii=False)
        os.replace(tmp, target)

        if cfg.copy_dashboard_files:
            source = os.path.join(BASE_DIR, "dashboard", "index.html")
            if os.path.isfile(source):
                shutil.copy2(source, os.path.join(cfg.web_root, "index.html"))
        log.info("Dashboard feed published to %s", target)
    except OSError as exc:
        log.error("Could not publish dashboard feed: %s", exc)


# ---------------------------------------------------------------------------
# Alerting
# ---------------------------------------------------------------------------
def load_alert_state(cfg: Config) -> dict:
    try:
        with open(cfg.alert_state_file, "r", encoding="utf-8") as handle:
            return json.load(handle)
    except (OSError, ValueError):
        return {}


def save_alert_state(cfg: Config, state: dict) -> None:
    try:
        os.makedirs(os.path.dirname(cfg.alert_state_file) or ".", exist_ok=True)
        with open(cfg.alert_state_file, "w", encoding="utf-8") as handle:
            json.dump(state, handle, indent=1)
    except OSError as exc:
        log.warning("Could not persist alert state: %s", exc)


def alert_key(result: dict) -> str:
    return "{}|{}|{}".format(
        result.get("server_name", ""), result.get("environment", ""), result.get("service_name", "")
    )


# Environments rendered with a heavier accent in the alert e-mail.
CRITICAL_ENVIRONMENTS = {"PROD", "PRODUCTION", "PRD", "LIVE"}


def environment_order(name: str) -> tuple:
    """Sort key so PROD-like environments always lead."""
    ranking = ["PROD", "PRODUCTION", "PRD", "LIVE", "PRE-PROD", "PREPROD", "DR",
               "UAT", "TEST", "QA", "DEV", "TRAINING"]
    upper = (name or "").upper()
    return (ranking.index(upper) if upper in ranking else len(ranking), upper)


def env_list(items: list) -> list:
    """Distinct environments in the batch, most critical first."""
    return sorted({(i.get("environment") or "-") for i in items}, key=environment_order)


def build_alert_html(cfg: Config, failures: list, run_id: str) -> str:
    def esc(value):
        return html.escape(str(value or ""))

    environments = env_list(failures)
    blocks = []
    for env in environments:
        group = [f for f in failures if (f.get("environment") or "-") == env]
        critical = env.upper() in CRITICAL_ENVIRONMENTS
        accent = "#d03b3b" if critical else "#1F3864"
        rows = []
        for f in group:
            rows.append(
                "<tr>"
                f"<td><b>{esc(f['server_name'])}</b></td>"
                f"<td>{esc(f['service_name'])}</td>"
                f"<td>{esc(f['failure_reason'])}</td>"
                f"<td style='text-align:center'>{esc(f['http_status'] or '-')}</td>"
                f"<td style='text-align:right'>{esc(f['response_ms'] or '-')}</td>"
                f"<td style='text-align:center'>{esc(f['attempts'])}</td>"
                f"<td style='word-break:break-all'>{esc(f['endpoint_url'])}</td>"
                f"<td>{esc(f['detail'])}</td>"
                "</tr>"
            )
        blocks.append(
            f"<p style='margin:18px 0 6px'><span style=\"background:{accent};color:#fff;"
            f"padding:3px 10px;border-radius:3px;font-weight:bold;letter-spacing:.5px\">"
            f"{esc(env)}</span> "
            f"<span style='color:#555'>&mdash; {len(group)} check(s) failed"
            f"{' &mdash; PRODUCTION' if critical else ''}</span></p>"
            "<table cellpadding='6' cellspacing='0' border='1' "
            "style='border-collapse:collapse;border-color:#ccc;font-size:12px;width:100%'>"
            "<thead style='background:#f3f3f1'><tr>"
            "<th align='left'>Server</th><th align='left'>Service</th><th align='left'>Reason</th>"
            "<th>HTTP</th><th>ms</th><th>Tries</th><th align='left'>Endpoint</th>"
            "<th align='left'>Detail</th></tr></thead>"
            f"<tbody>{''.join(rows)}</tbody></table>"
        )

    headline = ", ".join(environments)
    return f"""<html><body style="font-family:Segoe UI,Arial,sans-serif;font-size:13px;color:#222">
<p style="font-size:15px;margin-bottom:4px"><b style="color:#d03b3b">BSSV connectivity check FAILED</b>
&mdash; <b>{esc(headline)}</b></p>
<p style="margin-top:0">{len(failures)} check(s) across {len(environments)} environment(s) failed at
<b>{esc(iso_local(utcnow()))}</b> (run <code>{esc(run_id)}</code>) on host <b>{esc(platform.node())}</b>.</p>
{''.join(blocks)}
<p style="color:#666;font-size:11px;margin-top:20px">Sent by {APP_NAME} {APP_VERSION}. Successful checks are logged only, never mailed.
Alerts for the same server are throttled to one per {cfg.throttle_minutes} minute(s).</p>
</body></html>"""


def send_mail(cfg: Config, subject: str, body_html: str, to_addresses: str, attachments: list) -> bool:
    if not os.path.isfile(cfg.powershell_script):
        log.error("PowerShell mailer not found: %s", cfg.powershell_script)
        return False
    if not (cfg.smtp_server and cfg.from_address and to_addresses):
        log.error("E-mail not configured (smtp_server / from_address / to_addresses).")
        return False

    body_file = os.path.join(
        os.path.dirname(cfg.app_log) or BASE_DIR, f"alert_body_{uuid.uuid4().hex[:8]}.html"
    )
    try:
        with open(body_file, "w", encoding="utf-8") as handle:
            handle.write(body_html)

        cmd = [
            cfg.powershell_exe, "-NoProfile", "-NonInteractive",
            "-ExecutionPolicy", "Bypass",
            "-File", cfg.powershell_script,
            "-SmtpServer", cfg.smtp_server,
            "-Port", str(cfg.smtp_port),
            "-From", cfg.from_address,
            "-To", to_addresses,
            "-Subject", subject,
            "-BodyFile", body_file,
        ]
        if cfg.cc_addresses:
            cmd += ["-Cc", cfg.cc_addresses]
        if cfg.use_ssl:
            cmd += ["-UseSsl"]
        real_attachments = [a for a in attachments if a and os.path.isfile(a)]
        if cfg.attach_response and real_attachments:
            cmd += ["-Attachments", ";".join(real_attachments)]

        completed = subprocess.run(
            cmd, capture_output=True, text=True, timeout=120,
        )
        if completed.returncode == 0:
            log.info("Alert e-mail sent to %s", to_addresses)
            return True
        log.error(
            "Alert e-mail FAILED (exit %s): %s %s",
            completed.returncode,
            (completed.stdout or "").strip()[:500],
            (completed.stderr or "").strip()[:500],
        )
        return False
    except FileNotFoundError:
        log.error("PowerShell executable not found: %s", cfg.powershell_exe)
        return False
    except subprocess.TimeoutExpired:
        log.error("PowerShell mailer timed out after 120s")
        return False
    finally:
        try:
            os.remove(body_file)
        except OSError:
            pass


def dispatch_alerts(cfg: Config, results: list, run_id: str, dry_run: bool) -> None:
    if not cfg.email_enabled:
        return

    state = load_alert_state(cfg)
    now = utcnow()
    throttle = timedelta(minutes=max(0, cfg.throttle_minutes))

    to_alert, suppressed, recovered = [], 0, []
    for result in results:
        key = alert_key(result)
        entry = state.get(key, {})
        if result["status"] == "SUCCESS":
            if entry.get("alerted") and cfg.send_recovery:
                recovered.append(result)
            state[key] = {"alerted": False, "last_alert_utc": entry.get("last_alert_utc", "")}
            continue

        last = parse_utc(entry.get("last_alert_utc", ""))
        if throttle and last and (now - last) < throttle:
            suppressed += 1
            state[key] = entry
            continue
        to_alert.append(result)

    if suppressed:
        log.info("%s alert(s) suppressed by the %s-minute throttle", suppressed, cfg.throttle_minutes)

    if dry_run:
        if to_alert:
            log.warning("[DRY RUN] %s alert e-mail(s) would have been sent", len(to_alert))
        return

    def mark(items, sent):
        for item in items:
            item["email_sent"] = "YES" if sent else "ERROR"
            state[alert_key(item)] = {"alerted": True, "last_alert_utc": iso_utc(now)}

    if to_alert:
        if cfg.consolidate:
            # One mail per recipient list, and - by default - per environment, so
            # PROD alerts stay in their own thread and can carry their own rule.
            groups = {}
            for item in to_alert:
                key = (
                    item.get("_alert_to") or cfg.to_addresses,
                    (item.get("environment") or "-") if cfg.group_by_environment else "",
                )
                groups.setdefault(key, []).append(item)

            for key in sorted(groups, key=lambda k: environment_order(k[1])):
                items = groups[key]
                recipients = key[0]
                envs = ", ".join(env_list(items))
                names = ", ".join(sorted({i["server_name"] for i in items}))
                subject = (
                    f"{cfg.subject_prefix} {envs} - {len(items)} BSSV check(s) FAILED - {names}"
                )[:240]
                attachments = [i.get("_response_file", "") for i in items]
                sent = send_mail(cfg, subject, build_alert_html(cfg, items, run_id), recipients, attachments)
                mark(items, sent)
        else:
            for item in to_alert:
                subject = (
                    f"{cfg.subject_prefix} {item['environment']} - {item['server_name']} / "
                    f"{item['service_name']} FAILED - {item['failure_reason']}"
                )[:240]
                sent = send_mail(
                    cfg, subject, build_alert_html(cfg, [item], run_id),
                    item.get("_alert_to") or cfg.to_addresses,
                    [item.get("_response_file", "")],
                )
                mark([item], sent)

    for item in recovered:
        subject = (
            f"{cfg.subject_prefix} RECOVERED - {item['environment']} - "
            f"{item['server_name']} / {item['service_name']}"
        )
        body = (
            "<html><body style='font-family:Segoe UI,Arial,sans-serif'>"
            f"<p><b style='color:#0ca30c'>RECOVERED</b> &mdash; <b>{html.escape(item['environment'])}</b> "
            f"&mdash; {html.escape(item['server_name'])} / "
            f"{html.escape(item['service_name'])} responded successfully at "
            f"{html.escape(item['timestamp_local'])}.</p></body></html>"
        )
        send_mail(cfg, subject, body, item.get("_alert_to") or cfg.to_addresses, [])

    save_alert_state(cfg, state)


# ---------------------------------------------------------------------------
# One monitoring pass
# ---------------------------------------------------------------------------
def execute_pass(cfg: Config, dry_run: bool = False) -> dict:
    run_id = utcnow().strftime("%Y%m%d%H%M%S") + "-" + uuid.uuid4().hex[:4]
    try:
        servers = read_servers(cfg)
    except Exception as exc:  # noqa: BLE001
        log.error("Could not read the Excel server list: %s", exc)
        return {"run_id": run_id, "error": str(exc)}

    if not servers:
        log.warning("No enabled servers found in %s", cfg.excel_config)
        return {"run_id": run_id, "checked": 0, "success": 0, "failed": 0}

    log.info("Run %s - probing %s server(s)", run_id, len(servers))
    results = run_pass(cfg, servers, run_id)

    by_name = {(s["server_name"], s["environment"], s["service_name"]): s for s in servers}
    for result in results:
        server = by_name.get((result["server_name"], result["environment"], result["service_name"]))
        result["_alert_to"] = (server or {}).get("alert_email_to") or cfg.to_addresses
        if result["status"] != "SUCCESS":
            result["_response_file"] = save_failed_response(cfg, result)
            log.error(
                "FAILED  %s / %s / %s  (%s) %s",
                result["server_name"], result["environment"], result["service_name"],
                result["failure_reason"], result["detail"][:200],
            )
        else:
            log.info(
                "SUCCESS %s / %s / %s  %sms",
                result["server_name"], result["environment"], result["service_name"],
                result["response_ms"],
            )

    dispatch_alerts(cfg, results, run_id, dry_run)

    append_results(cfg, results)
    purge_log(cfg)
    purge_response_files(cfg)

    if cfg.dashboard_enabled:
        publish_dashboard(cfg, build_status_json(cfg, read_log(cfg)))

    success = sum(1 for r in results if r["status"] == "SUCCESS")
    summary = {
        "run_id": run_id,
        "checked": len(results),
        "success": success,
        "failed": len(results) - success,
    }
    log.info(
        "Run %s complete - %s checked, %s success, %s failed",
        run_id, summary["checked"], summary["success"], summary["failed"],
    )
    return summary


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------
def main() -> int:
    parser = argparse.ArgumentParser(description=f"{APP_NAME} {APP_VERSION}")
    parser.add_argument("--config", default=os.path.join(BASE_DIR, "config.ini"),
                        help="path to config.ini (default: alongside this script)")
    parser.add_argument("--once", action="store_true", help="run a single pass and exit")
    parser.add_argument("--loop", action="store_true", help="stay resident and repeat")
    parser.add_argument("--dry-run", action="store_true", help="probe and log, but never send e-mail")
    parser.add_argument("--test-email", action="store_true", help="send a test alert and exit")
    parser.add_argument("--verbose", action="store_true", help="debug logging on the console")
    args = parser.parse_args()

    cfg = Config(resolve(args.config, os.getcwd()))
    setup_logging(cfg, args.verbose)

    if not cfg.verify_ssl:
        urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
        log.debug("TLS verification disabled (verify_ssl = false)")

    log.info("%s %s starting - config: %s", APP_NAME, APP_VERSION, cfg.path)

    if args.test_email:
        sample = {
            "server_name": "TEST-SERVER", "environment": "TEST", "service_name": "RI_AddressBookManager",
            "endpoint_url": "https://example.invalid/PD920/RI_AddressBookManager",
            "failure_reason": "TEST_ALERT", "detail": "This is a test message from BSSV Monitor.",
            "http_status": "000", "response_ms": "0", "attempts": 1,
        }
        ok = send_mail(
            cfg, f"{cfg.subject_prefix} Test alert", build_alert_html(cfg, [sample], "test"),
            cfg.to_addresses, [],
        )
        print("Test e-mail sent." if ok else "Test e-mail FAILED - see the application log.")
        return 0 if ok else 1

    once = args.once or (cfg.run_mode == "once" and not args.loop)
    if once:
        summary = execute_pass(cfg, args.dry_run)
        return 1 if summary.get("error") else 0

    interval = max(10, cfg.run_interval_seconds)
    log.info("Loop mode - a pass every %s second(s). Ctrl+C to stop.", interval)
    next_run = time.monotonic()
    try:
        while True:
            execute_pass(cfg, args.dry_run)
            next_run += interval
            now = time.monotonic()
            if next_run <= now:
                if cfg.skip_missed_ticks:
                    missed = int((now - next_run) // interval) + 1
                    next_run += missed * interval
                    log.warning("Pass overran the interval - skipping %s tick(s)", missed)
                else:
                    next_run = now
            time.sleep(max(0, next_run - time.monotonic()))
    except KeyboardInterrupt:
        log.info("Stopped by operator.")
        return 0


if __name__ == "__main__":
    sys.exit(main())
